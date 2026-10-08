import sys, json, copy, statistics as st
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

S = sys.argv[1]
TEMPLATE, OUT = sys.argv[2], sys.argv[3]
D = json.load(open(f'{S}/work/data.json'))

# ---------- palette (fra DMP-malen) ----------
DARK = RGBColor(0x28, 0x30, 0x17)
GREEN = RGBColor(0xBA, 0xF6, 0xAE)
CREAM = RGBColor(0xFC, 0xEB, 0xC8)
LAV = RGBColor(0xCD, 0xCE, 0xFB)
CYAN = RGBColor(0xCB, 0xF9, 0xFE)
PINK = RGBColor(0xF6, 0xC8, 0xFD)
OLIVE = RGBColor(0x5B, 0x5B, 0x2B)
ORANGE = RGBColor(0xF2, 0x6B, 0x43)
PURPLE = RGBColor(0x8A, 0x2B, 0xC2)
LIGHTOLIVE = RGBColor(0x9C, 0xA3, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x6B, 0x70, 0x5C)

STATUS = {  # fill, text
    'Positiv': (GREEN, DARK),
    'Blandet': (LAV, DARK),
    'Uendret': (CREAM, DARK),
    'Negativ': (ORANGE, DARK),
    'Ikke dokumentert': (GREY, WHITE),
}


def f1(x):
    return f'{x:.1f}'.replace('.', ',')


def pct(x):
    return f'{round(x * 100):d} %'


prs = Presentation(TEMPLATE)
# fjern eksisterende lysbilder
sldIdLst = prs.slides._sldIdLst
for s in list(sldIdLst):
    prs.part.drop_rel(s.rId)
    sldIdLst.remove(s)
L = {l.name: l for l in prs.slide_layouts}


def clear_unused_placeholders(slide, keep=()):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in keep:
            ph._element.getparent().remove(ph._element)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def tb(slide, x, y, w, h, paras, size=12, color=DARK, anchor=MSO_ANCHOR.TOP, margin=0.05, name=None):
    """paras: list av str eller dict(text|runs, size, bold, color, bullet, align, space_after, italic)"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', Inches(margin))
    first = True
    for p in paras:
        if isinstance(p, str):
            p = {'text': p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        runs = p.get('runs') or [(p.get('text', ''), {})]
        for txt, ro in runs:
            r = para.add_run()
            r.text = txt
            r.font.size = Pt(ro.get('size', p.get('size', size)))
            r.font.bold = ro.get('bold', p.get('bold', False))
            r.font.italic = ro.get('italic', p.get('italic', False))
            r.font.color.rgb = ro.get('color', p.get('color', color))
        if p.get('align'):
            para.alignment = p['align']
        para.space_after = Pt(p.get('space_after', 4))
        if p.get('bullet'):
            pPr = para._p.get_or_add_pPr()
            lvl = p.get('level', 0)
            pPr.set('marL', str(int(Inches(0.18 + 0.18 * lvl))))
            pPr.set('indent', str(int(-Inches(0.16))))
            bu = etree.SubElement(pPr, qn('a:buChar'))
            bu.set('char', '•' if lvl == 0 else '–')
    return box


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ''
    return s


def chip(slide, x, y, w, h, text, size=9):
    fill, fg = STATUS[text]
    s = rect(slide, x, y, w, h, fill, radius=0.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = fg
    return s


def content_slide(title, subtitle=None):
    s = prs.slides.add_slide(L['Tekst og innhold'])
    clear_unused_placeholders(s, keep=(0,))
    s.shapes.title.text = title
    for p in s.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(20)
    if subtitle:
        tb(s, 0.47, 0.88, 9.06, 0.3, [{'text': subtitle, 'size': 11, 'color': OLIVE, 'italic': True}], margin=0)
    return s


def dark_slide(title, subtitle=None):
    s = prs.slides.add_slide(L['Fullsidebilde'])
    clear_unused_placeholders(s)
    tb(s, 0.47, 0.3, 9.06, 0.45, [{'text': title, 'size': 20, 'color': GREEN}], margin=0)
    if subtitle:
        tb(s, 0.47, 0.75, 9.06, 0.3, [{'text': subtitle, 'size': 11, 'color': CREAM, 'italic': True}], margin=0)
    return s


def section(layout, title, sub):
    s = prs.slides.add_slide(L[layout])
    s.shapes.title.text = title
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = sub
    return s


def transparent_chart(chart):
    cs = chart._chartSpace
    for old in cs.findall(qn('c:spPr')):
        cs.remove(old)
    sp = etree.SubElement(cs, qn('c:spPr'))
    etree.SubElement(sp, qn('a:noFill'))
    ln = etree.SubElement(sp, qn('a:ln'))
    etree.SubElement(ln, qn('a:noFill'))
    # spPr must come right after c:chart
    cs.remove(sp)
    cs.find(qn('c:chart')).addnext(sp)


def style_axes(chart, vmin=1, vmax=6, show_val=True, size=9, color=DARK):
    va = chart.value_axis
    va.minimum_scale = vmin
    va.maximum_scale = vmax
    va.major_unit = 1
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LIGHTOLIVE
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size)
    va.tick_labels.font.color.rgb = color
    va.visible = show_val
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(size)
    ca.tick_labels.font.color.rgb = color
    ca.format.line.color.rgb = LIGHTOLIVE
    ca.has_major_gridlines = False


def legend(chart, size=9, color=DARK, pos=XL_LEGEND_POSITION.BOTTOM):
    chart.has_legend = True
    chart.legend.position = pos
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(size)
    chart.legend.font.color.rgb = color


def line_chart(slide, x, y, w, h, cats, series, colors, title=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, [round(v, 1) for v in vals])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    transparent_chart(ch)
    ch.font.size = Pt(9)
    ch.font.color.rgb = DARK
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = DARK
    else:
        ch.has_title = False
    style_axes(ch)
    for s, c in zip(ch.series, colors):
        s.smooth = False
        s.format.line.color.rgb = c
        s.format.line.width = Pt(2.25)
        s.marker.style = XL_MARKER_STYLE.CIRCLE
        s.marker.size = 6
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = c
        s.marker.format.line.color.rgb = c
    # tallverdi bare på siste punkt (2026)
    for sr, pos in zip(ch.series, [XL_LABEL_POSITION.RIGHT, XL_LABEL_POSITION.BELOW, XL_LABEL_POSITION.ABOVE]):
        dlb = sr.points[len(cats) - 1].data_label
        dlb.has_text_frame = False
        dlb.show_value = True
        dlb.number_format = '0.0'
        dlb.number_format_is_linked = False
        dlb.position = pos
        dlb.font.size = Pt(8)
        dlb.font.bold = True
        dlb.font.color.rgb = DARK
    ch.has_legend = False
    return ch


def bar_chart(slide, x, y, w, h, cats, series, colors, horizontal=True, labels=True, vmin=1, vmax=6,
              stacked=False, gap=60, overlap=None, size=9, label_color=DARK, legend_on=True, numfmt='0.0'):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    if stacked:
        t = XL_CHART_TYPE.BAR_STACKED_100 if horizontal else XL_CHART_TYPE.COLUMN_STACKED_100
    else:
        t = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(t, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    transparent_chart(ch)
    ch.font.size = Pt(size)
    ch.font.color.rgb = DARK
    ch.has_title = False
    if stacked:
        va = ch.value_axis
        va.visible = False
        va.has_major_gridlines = False
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(size)
        ca.tick_labels.font.color.rgb = DARK
        ca.format.line.fill.background()
    else:
        style_axes(ch, vmin, vmax, show_val=not horizontal, size=size)
        if horizontal:
            ch.value_axis.has_major_gridlines = False
    if horizontal:
        ch.category_axis.reverse_order = True
    pl = ch.plots[0]
    pl.gap_width = gap
    if overlap is not None:
        pl.overlap = overlap
    for s, c in zip(ch.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c
        s.invert_if_negative = False
    if labels:
        pl.has_data_labels = True
        dl = pl.data_labels
        dl.number_format = numfmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(size - 1)
        dl.font.color.rgb = label_color
        dl.position = XL_LABEL_POSITION.CENTER if stacked else XL_LABEL_POSITION.OUTSIDE_END
        if stacked:
            for sr, c in zip(ch.series, colors):
                sdl = sr.data_labels
                sdl.number_format = numfmt
                sdl.number_format_is_linked = False
                sdl.show_value = True
                sdl.position = XL_LABEL_POSITION.CENTER
                sdl.font.size = Pt(size - 1)
                sdl.font.color.rgb = CREAM if c == DARK else DARK
    if legend_on:
        legend(ch, size=size - 1)
    else:
        ch.has_legend = False
    return ch


def table(slide, x, y, w, col_w, rows, row_h=0.22, size=8, header_fill=GREEN, header_fg=DARK,
          body_fill=None, body_fg=CREAM, fills=None, bold_rows=(), aligns=None, fgs=None):
    nr, nc = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = gf.table
    tblPr = t._tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i in range(nr):
        t.rows[i].height = Inches(row_h)
        for j in range(nc):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.05)
            c.margin_top = c.margin_bottom = Inches(0.015)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = header_fill if i == 0 else (fills[i][j] if fills and fills[i][j] is not None else body_fill)
            if fill is None:
                c.fill.background()
            else:
                c.fill.solid()
                c.fill.fore_color.rgb = fill
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(rows[i][j])
            r.font.size = Pt(size)
            r.font.bold = (i == 0) or (i in bold_rows)
            fg = header_fg if i == 0 else body_fg
            if fgs and i > 0 and fgs[i][j] is not None:
                fg = fgs[i][j]
            r.font.color.rgb = fg
            if aligns:
                p.alignment = aligns[j]
    return t


# =====================================================================
# Data
R1 = D['R1']; R2 = D['R2']
F = ['Apotekkonsesjoner', 'Tilsyn', 'Virksomhetstillatelser']
FS = {'Apotekkonsesjoner': 'Apotek', 'Tilsyn': 'Tilsyn', 'Virksomhetstillatelser': 'Virksomhetstillatelser'}
Q1 = R1['q']; Q2 = R2['q']
R1MAP = [1, 2, 3, 5, 6, 7, 8, 9, 10, 11]  # R1-spm -> R2-indeks
a25 = R2['vals']['2025']['Alle']; a26 = R2['vals']['2026']['Alle']
mi1 = R1['mi']; mi2 = R2['mi']

# ledere samlet R1+R2 2026
def load(name):
    ws = openpyxl.load_workbook(f'{S}/src/{name}.xlsx').worksheets[0]
    return list(ws.iter_rows(values_only=True))[1:]
def num(v):
    try: return int(v)
    except: return None
r1raw = load('d32ea74a-R1oktober2026'); r2raw = load('0025ce6a-R2oktober2026')
pooled = [(r[7].lower(), [num(r[q]) for q in range(8, 18)]) for r in r1raw] + \
         [(r[7].lower(), [num(r[8 + k]) for k in R1MAP]) for r in r2raw]
def role_means(role):
    sub = [v for rl, v in pooled if rl == role]
    return len(sub), [st.mean([x[i] for x in sub if x[i]]) for i in range(10)]
nL, mL = role_means('enhetsleder')
nS, mS = role_means('saksbehandler/lagleder')

# =====================================================================
# 1 Forside
s = prs.slides.add_slide(L['Tittellysbilde'])
s.shapes.title.text = 'Gevinstrealisering i DELE'
for ph in s.placeholders:
    if ph.placeholder_format.idx == 1:
        ph.text = 'Gevinstmålinger for R1 og R2\nRapport til prosjekteier – oktober 2026'
    elif ph.placeholder_format.idx not in (0, 1):
        ph._element.getparent().remove(ph._element)
notes(s, 'Rapporten oppsummerer gevinstmålingene for DELE R1 (nullpunkt sep. 2022, puls mai 2023, ny måling okt. 2026) '
         'og R2 (nullpunkt sommer 2025, ny måling okt. 2026), sett opp mot vedtatt gevinstrealiseringsplan.')

# =====================================================================
# 2 Hovedbudskap
s = content_slide('Hovedbudskap')
cards = [
    ('R1: Gevinstene er ikke realisert i Tilsyn',
     'Tilsyn (14 av 18 svar) ligger på eller under nullpunktet fra 2022 for alle tre måleindikatorer. '
     'Virksomhetstillatelser har bedret seg tydelig. Apotek ligger under nullpunktet, men medarbeidertilfredsheten har tatt seg opp siden 2023.'),
    ('R2: Fremgang på styring og oversikt',
     'Oversikt over frister, prioritering/ressursstyring og intern samhandling har økt med ca. ½ poeng, og økningen er statistisk sikker. '
     'Kvalitet, ekstern samhandling og informasjon til eksterne er uendret.'),
    ('Brukervennlighet og parallelle systemer bremser',
     'Andelen som er uenig i at verktøyene er enkle å bruke, økte fra 1 % til 17 % i R2. '
     'Fritekstsvarene peker på mange klikk, at ikke alle saker ligger i DELE og at mye arbeid fortsatt skjer i Excel, Teams og e-post.'),
    ('To gevinster mangler dokumentasjon',
     'Tidsbesparelse (R1 pri 2.1) og redusert forvaltningskostnad for SAM-T (R1 pri 1) er ikke målt. '
     'Ingen gevinster har målverdier, så rapporten viser retning, ikke måloppnåelse.'),
]
for i, (h, t) in enumerate(cards):
    cx = 0.47 + (i % 2) * 4.6
    cy = 1.3 + (i // 2) * 1.85
    rect(s, cx, cy, 4.4, 1.7, CREAM if i != 2 else LAV, name=f'Kort {i+1}')
    tb(s, cx + 0.15, cy + 0.1, 0.4, 0.45, [{'text': str(i + 1), 'size': 22, 'bold': True, 'color': ORANGE}], margin=0)
    tb(s, cx + 0.55, cy + 0.12, 3.7, 1.5, [{'text': h, 'size': 12, 'bold': True, 'space_after': 4},
                                         {'text': t, 'size': 10}], margin=0)
notes(s, 'Fire hovedbudskap. 1) R1 er i praksis et Tilsyn-resultat fordi 14 av 18 svar kommer derfra. '
         '2) R2 viser forbedring på nettopp det DELE skal løse: oversikt, prioritering og samhandling internt (p<0,05). '
         '3) Spredningen i R2 øker: flere er fornøyde, men også flere misfornøyde. 4) Gevinstplanen har ingen målverdier og to indikatorer er ikke målt.')

# =====================================================================
# 3 Gevinststatus (mørk tabell)
s = dark_slide('Gevinststatus – oversikt', 'Utvikling fra nullpunkt til okt. 2026 (snitt, skala 1–6). Ingen målverdier er fastsatt, så status viser retning.')
rows = [['Rel.', 'Gevinst (prioritet)', 'Måleindikator', 'Utvikling fra nullpunkt til 2026', 'Retning']]
stat = []
def tri(m):
    return ' | '.join(f'{FS[f][:3] if f!="Virksomhetstillatelser" else "VT"} {f1(mi1["2022"][f][m])}→{f1(mi1["2026"][f][m])}' for f in F)
rows += [
    ['R1', 'Pri 1: Forutsigbarhet i IT-utvikling og -kostnader', 'Vedlikeholdskostnad SAM-T (161 000 kr/md.)', 'Ikke målt i dette grunnlaget', 'Ikke dokumentert'],
    ['R1', 'Pri 2: Forbedret intern saksbehandling', '2.1 Tid til koordinering, ressursfordeling, rapportering', 'Tidsmåling er ikke gjennomført', 'Ikke dokumentert'],
    ['R1', 'Pri 2: Forbedret intern saksbehandling', '2.2 Medarbeidertilfredshet', tri('Medarbeidertilfredshet'), 'Blandet'],
    ['R1', 'Pri 3: Økt kvalitet på faglige vurderinger', 'Kvalitet på faglige vurderinger', tri('Kvalitet'), 'Negativ'],
    ['R1', 'Pri 4: Økt samhandling med eksterne, nasjonale', 'Samhandling eksterne, nasjonale', tri('Samhandling nasjonal'), 'Negativ'],
]
m25, m26 = mi2['2025']['Alle'], mi2['2026']['Alle']
def d2(m, extra=''):
    return f'{f1(m25[m])} → {f1(m26[m])}{extra}'
rows += [
    ['R2', 'Forbedret intern saksbehandling', 'Medarbeidertilfredshet', d2('Medarbeidertilfredshet', '  (frister, prioritering og intern samhandling ↑)'), 'Positiv'],
    ['R2', 'Økt kvalitet på faglige vurderinger', 'Kvalitet på faglige vurderinger', d2('Kvalitet'), 'Uendret'],
    ['R2', 'Økt samhandling med eksterne, nasjonale', 'Samhandling eksterne, nasjonale', d2('Samhandling nasjonal'), 'Uendret'],
    ['R2', 'Økt samhandling med eksterne, internasjonale', 'Samhandling eksterne, internasjonale', d2('Samhandling internasjonal'), 'Uendret'],
    ['R2', 'Informasjon til eksterne enklere tilgjengelig', 'Tilgjengelighet av info for eksterne', d2('Informasjonstilgjengelighet', '  (41 % svarer «ikke relevant»)'), 'Uendret'],
]
statuses = [r[4] for r in rows[1:]]
for r in rows[1:]:
    r[4] = ''
cw = [0.4, 2.45, 2.35, 2.75, 1.1]
t = table(s, 0.47, 1.15, sum(cw), cw, rows, row_h=0.33, size=8, body_fill=None)
for i, stx in enumerate(statuses):
    chip(s, 0.47 + sum(cw[:4]) + 0.08, 1.15 + 0.33 * (i + 1) + 0.06, 0.95, 0.21, stx, size=7)
tb(s, 0.47, 4.86, 9.06, 0.25, [{'text': 'Apo = Apotekkonsesjoner, Til = Tilsyn, VT = Virksomhetstillatelser. R1 2026: Apotek n=3, VT n=4 (tre personer svarte for begge), Tilsyn n=14.',
                              'size': 8, 'color': CREAM, 'italic': True}], margin=0)
notes(s, 'Statusen er en faglig vurdering av retning, siden det ikke er fastsatt målverdier. '
         '«Blandet» for R1 medarbeidertilfredshet: Tilsyn ned (3,2→2,5), Apotek ned fra et svært høyt nullpunkt (5,2→4,0) men opp fra 2023 (3,4), VT opp (2,9→4,1). '
         'For R2 er medarbeidertilfredshet den eneste indikatoren med tydelig fremgang. Den skyldes særlig tre spørsmål med statistisk sikker økning.')

# =====================================================================
# 4 Datagrunnlag
s = content_slide('Datagrunnlag og forbehold')
meas = [
    ('R1', 'Nullpunkt', 'sep. 2022', 'Snitt per fagområde', 'n ukjent'),
    ('R1', 'Puls', 'mai 2023', 'Snitt per fagområde', 'n ukjent'),
    ('R1', 'Ny måling', 'okt. 2026', '18 svar av ca. 30', '≈ 60 %'),
    ('R2', 'Nullpunkt', 'sommer 2025', f'{R2["n"]["2025"]} svar av ca. 150', '≈ 50 %'),
    ('R2', 'Ny måling', 'okt. 2026', f'{R2["n"]["2026"]} svar av ca. 150', '≈ 39 %'),
]
for i, (rel, typ, when, n, rr) in enumerate(meas):
    x = 0.47 + i * 1.84
    rect(s, x, 1.25, 1.7, 1.45, LAV if rel == 'R1' else CYAN, name=f'Måling {i+1}')
    tb(s, x + 0.1, 1.32, 1.5, 1.35, [
        {'text': f'{rel} – {typ}', 'size': 10, 'bold': True},
        {'text': when, 'size': 16, 'bold': True, 'color': DARK, 'space_after': 2},
        {'text': n, 'size': 9},
        {'text': f'Svarprosent {rr}' if '%' in rr else rr, 'size': 9, 'color': OLIVE},
    ], margin=0)
tb(s, 0.47, 2.9, 4.4, 2.3, [
    {'text': 'Slik er tallene beregnet', 'size': 12, 'bold': True},
    {'text': 'Skala 1 (helt uenig) – 6 (helt enig). «Ikke relevant» er holdt utenfor snittet.', 'size': 10, 'bullet': True},
    {'text': 'Måleindikator = snitt av spørsmålene som hører til, som i gevinstplanen.', 'size': 10, 'bullet': True},
    {'text': 'R1 2026 er regnet om til snitt per fagområde for å kunne sammenlignes med 2022 og 2023. Den som svarte for flere fagområder, er telt i hvert av dem.', 'size': 10, 'bullet': True},
    {'text': 'R2: «Regulatorisk PI» (2025) = «Regulatorisk produktinformasjon» (2026). Små fagområder er slått sammen.', 'size': 10, 'bullet': True},
], margin=0)
tb(s, 5.1, 2.9, 4.4, 2.3, [
    {'text': 'Forbehold', 'size': 12, 'bold': True},
    {'text': 'Det er få svar per gruppe: R1 Apotek har 3 og Virksomhetstillatelser 4. Endringer under ca. 0,5 poeng bør ikke tillegges vekt.', 'size': 10, 'bullet': True},
    {'text': 'Undersøkelsene er anonyme, så vi kan ikke følge samme personer over tid. Lavere svarprosent i R2 2026 kan påvirke resultatet.', 'size': 10, 'bullet': True},
    {'text': 'R1 2026 er målt fire år etter innføring og etter R2. Andre endringer (organisering, nye verktøy) kan påvirke svarene.', 'size': 10, 'bullet': True},
    {'text': 'Ingen målverdier er fastsatt, så rapporten viser utvikling og ikke måloppnåelse.', 'size': 10, 'bullet': True},
], margin=0)
notes(s, 'Statistisk test: For R2 er forskjellen mellom 2025 og 2026 testet med en permutasjonstest på snittet (p<0,05 regnes som sikker). '
         'For R1 finnes bare snitt for 2022 og 2023, så det er ikke mulig å teste der.')

# =====================================================================
# 5 Section R1
s = section('Deloverskrift Piller', 'Release 1', 'Apotekkonsesjoner, Tilsyn og Virksomhetstillatelser\nGo-live oktober 2022')
notes(s, 'Resultater for R1.')

# 6 R1 utvikling per måleindikator
s = content_slide('R1: Bare Virksomhetstillatelser ligger over nullpunktet i 2026',
                  'Måleindikatorer per fagområde – snitt, skala 1–6. n = antall svar i 2026')
mis = ['Kvalitet', 'Samhandling nasjonal', 'Medarbeidertilfredshet']
milab = ['Kvalitet', 'Ekstern nasjonal samhandling', 'Medarbeidertilfredshet']
years = ['Sep. 2022', 'Mai 2023', 'Okt. 2026']
for k, (lab, c, lx) in enumerate(zip(milab, [DARK, ORANGE, PURPLE], [0.5, 1.5, 3.65])):
    rect(s, lx, 1.27, 0.13, 0.13, c, shape=MSO_SHAPE.OVAL)
    tb(s, lx + 0.2, 1.2, 2.2, 0.28, [{'text': lab, 'size': 9}], margin=0)
for i, f in enumerate(F):
    series = [(milab[k], [mi1[y][f][m] for y in ('2022', '2023', '2026')]) for k, m in enumerate(mis)]
    line_chart(s, 0.35 + i * 3.1, 1.5, 3.1, 3.25, years, series, [DARK, ORANGE, PURPLE],
               title=f'{FS[f]} (n = {R1["n2026"][f]})')
tb(s, 0.47, 4.8, 9.06, 0.5, [{'text': 'Apotek hadde et svært høyt nullpunkt (5,2–5,8), og det gir lite rom for forbedring. Tilsyn har det laveste nivået og ingen bedring siden 2023. '
                             'Virksomhetstillatelser har bedret seg på alle tre indikatorer, mest på medarbeidertilfredshet (2,9 → 4,1).', 'size': 10}], margin=0)
notes(s, 'Kurvene viser det typiske fallet like etter innføring (2023). Spørsmålet er om nivået har tatt seg opp igjen. '
         'VT: ja, over nullpunkt. Apotek: delvis, medarbeidertilfredshet opp fra 3,4 til 4,0, men kvalitet og ekstern samhandling har fortsatt å falle. '
         'Tilsyn: nei, alle tre er på eller under 2023-nivå, og ekstern samhandling er lavest (1,6). '
         'NB: Apotek og VT bygger på 3 og 4 svar, og tre av dem er de samme personene.')

# 7 R1 tabell
s = dark_slide('R1: Utvikling per spørsmål og fagområde', 'Snitt (1–6). Farge i 2026-kolonnen: grønn = minst 0,3 over nullpunkt, oransje = minst 0,3 under.')
hdr = ['Spørsmål']
for f in F:
    hdr += [f'{FS[f][:3] if f != "Virksomhetstillatelser" else "VT"} 22', '23', '26']
rows = [hdr]; fills = [[None] * 10]; fgs = [[None] * 10]
for qi, q in enumerate(Q1):
    row = [q]; fr = [None]; fg = [None]
    for f in F:
        v = [R1['vals'][y][f][qi] for y in ('2022', '2023', '2026')]
        row += [f1(x) for x in v]
        d = v[2] - v[0]
        fr += [None, None, GREEN if d >= 0.3 else (ORANGE if d <= -0.3 else CREAM)]
        fg += [None, None, DARK]
    rows.append(row); fills.append(fr); fgs.append(fg)
# MI-rader
for m, lab in zip(mis, ['MI Kvalitet', 'MI Ekstern nasjonal samhandling', 'MI Medarbeidertilfredshet']):
    row = [lab]; fr = [None]; fg = [None]
    for f in F:
        v = [mi1[y][f][m] for y in ('2022', '2023', '2026')]
        row += [f1(x) for x in v]
        d = v[2] - v[0]
        fr += [None, None, GREEN if d >= 0.3 else (ORANGE if d <= -0.3 else CREAM)]
        fg += [None, None, DARK]
    rows.append(row); fills.append(fr); fgs.append(fg)
cw = [2.7] + [0.7] * 9
table(s, 0.47, 1.15, sum(cw), cw, rows, row_h=0.28, size=8, body_fill=None, fills=fills, fgs=fgs,
      bold_rows=(11, 12, 13), aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 9)
notes(s, 'Detaljtabell for R1. Tilsyn ligger under nullpunktet på alle ti spørsmål. Størst fall: intern samhandling (4,0→2,4), faglig støtte (3,3→2,3) og status (3,9→2,9). '
         'VT: størst fremgang på å arbeide effektivt (2,2→4,8), status i saksbehandlingen (3,0→5,3) og oversikt over oppgaver (4,0→5,5).')

# 8 Tilsyn
s = content_slide('R1 Tilsyn: Svakest på frister, prioritering og ekstern samhandling',
                  f'Tilsyn, snitt 2022 og 2026 (n 2026 = {R1["n2026"]["Tilsyn"]})')
ti = [R1['vals']['2022']['Tilsyn'], R1['vals']['2026']['Tilsyn']]
bar_chart(s, 0.3, 1.15, 5.4, 4.2, Q1, [('Sep. 2022', ti[0]), ('Okt. 2026', ti[1])], [LIGHTOLIVE, DARK], vmin=0, vmax=6, gap=40, size=9)
quotes = [
    ('«Funksjonaliteten i DELE forandrer seg stadig og systemet oppleves som upålitelig. Må holde oversikt over frister på egen hånd utenfor DELE.»', 'Saksbehandler, Tilsyn'),
    ('«En del av spørsmålene indikerer at det skal være en viss arbeidsflyt … Det er det ikke – det ble ikke prioritert da DELE ble laget for oss.»', 'Saksbehandler, Tilsyn'),
    ('«PowerBI kunne vært koblet til, bedre dashbord.»', 'Enhetsleder, Tilsyn'),
]
yy = 1.2
for q, who in quotes:
    h = 1.25 if len(q) > 100 else 0.7
    rect(s, 5.9, yy, 3.65, h, CREAM)
    tb(s, 6.0, yy + 0.05, 3.45, h - 0.1, [{'text': q, 'size': 9, 'italic': True, 'space_after': 2},
                                         {'text': who, 'size': 8, 'color': OLIVE}], margin=0, anchor=MSO_ANCHOR.MIDDLE)
    yy += h + 0.15
tb(s, 5.9, yy, 3.65, 0.6, [{'text': 'Alle fire fritekstsvar i R1 kommer fra Tilsyn. De som skrev kommentar, har lavere snitt (2,2) enn resten (2,9).', 'size': 9}], margin=0)
notes(s, 'Tilsyn utgjør 14 av 18 svar i R1 2026, og R1-resultatet er derfor i hovedsak Tilsyns erfaring. '
         'Alle ti spørsmål ligger under nullpunktet fra 2022. Oversikt over frister bedret seg midlertidig i 2023 (2,9), men har falt tilbake (1,9). '
         'Fritekst og tall sier det samme: inspeksjoner mangler arbeidsflytstøtte, frister følges opp utenfor DELE, arkivet (360) er tungt å bruke, og ledere savner dashbord.')

# 9 Apotek og VT
s = content_slide('R1 Apotek og Virksomhetstillatelser: Bedre enn i 2023, men få svar',
                  'Snitt for utvalgte spørsmål – tre av svarene gjelder begge fagområder')
sel = [3, 4, 9, 6]  # status, effektivt, oppgaver, prioritering
for k, f in enumerate(['Apotekkonsesjoner', 'Virksomhetstillatelser']):
    x = 0.47 + k * 4.6
    rect(s, x, 1.25, 4.4, 2.8, CREAM if k == 0 else CYAN, name=f'Kort {FS[f]}')
    tb(s, x + 0.15, 1.33, 4.1, 0.35, [{'text': f'{FS[f]} (n = {R1["n2026"][f]})', 'size': 13, 'bold': True}], margin=0)
    rws = [['Spørsmål', '2022', '2023', '2026']]
    for qi in sel:
        rws.append([Q1[qi]] + [f1(R1['vals'][y][f][qi]) for y in ('2022', '2023', '2026')])
    table(s, x + 0.15, 1.75, 4.1, [2.3, 0.6, 0.6, 0.6], rws, row_h=0.3, size=9, header_fill=DARK, header_fg=CREAM,
          body_fill=None, body_fg=DARK, aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 3)
    txt = ('Under nullpunktet på alle 10 spørsmål, men bedre enn i 2023 på 6 av 10. Nullpunktet (4–6) var uvanlig høyt.' if k == 0 else
           'Over nullpunktet på alle 10 spørsmål. Størst fremgang på å arbeide effektivt, status og oversikt over oppgaver.')
    tb(s, x + 0.15, 3.4, 4.1, 1.1, [{'text': txt, 'size': 10}], margin=0)
tb(s, 0.47, 4.3, 9.06, 0.6, [{'text': 'Begge fagområder er mest positive til status, oversikt og effektivitet. Svakest er prioritering/ressursstyring og frister, som hos Tilsyn. '
                             'Med 3–4 svar kan resultatene snu ved neste måling.', 'size': 10}], margin=0)
notes(s, 'Apotek: 3 svar, alle har også krysset av for VT. Tolk forsiktig. Det høye nullpunktet i 2022 (snitt 5,2–5,8) kan skyldes få svar eller høye forventninger før innføring.')

# =====================================================================
# 10 Section R2
s = section('Deloverskrift Fisk', 'Release 2', 'Regulatoriske og kliniske fagområder\nGo-live september 2025')
notes(s, 'Resultater for R2.')

# 11 R2 per spørsmål
s = content_slide('R2: Fremgang på frister, prioritering og intern samhandling',
                  f'Snitt per spørsmål, alle svar (n = {R2["n"]["2025"]} i 2025 og {R2["n"]["2026"]} i 2026). * = statistisk sikker endring (p < 0,05)')
order = sorted(range(12), key=lambda i: -(a26[i]['m'] - a25[i]['m']))
cats = [Q2[i] + (' *' if R2['p'][i] < 0.05 else '') for i in order]
bar_chart(s, 0.3, 1.15, 6.0, 4.25, cats, [('Sommer 2025', [a25[i]['m'] for i in order]), ('Okt. 2026', [a26[i]['m'] for i in order])],
          [LIGHTOLIVE, DARK], vmin=0, vmax=6, gap=35, size=9)
rect(s, 6.5, 1.25, 3.05, 1.9, GREEN if False else CREAM)
tb(s, 6.6, 1.3, 2.85, 1.8, [
    {'text': 'Det som har bedret seg', 'size': 11, 'bold': True},
    {'text': 'Oversikt over frister +0,6', 'size': 10, 'bullet': True},
    {'text': 'Prioritering og ressursstyring +0,5', 'size': 10, 'bullet': True},
    {'text': 'Intern samhandling +0,4', 'size': 10, 'bullet': True},
    {'text': 'Dette er kjernen i det DELE skal gi: styring og oversikt.', 'size': 9, 'italic': True},
], margin=0)
rect(s, 6.5, 3.3, 3.05, 2.0, LAV)
tb(s, 6.6, 3.35, 2.85, 1.9, [
    {'text': 'Det som står stille eller faller', 'size': 11, 'bold': True},
    {'text': 'Enkle å bruke −0,3 (4,3 → 4,0)', 'size': 10, 'bullet': True},
    {'text': 'Info til eksterne −0,2', 'size': 10, 'bullet': True},
    {'text': 'Faglig støtte, finne faglig info, effektivitet og ekstern samhandling er uendret (±0,1).', 'size': 10, 'bullet': True},
], margin=0)
notes(s, 'Permutasjonstest på snitt: frister p=0,02, prioritering p=0,02, intern samhandling p=0,04. Enkle å bruke p=0,14 (ikke sikker, men retningen samsvarer med fritekst). '
         'De øvrige endringene er innenfor tilfeldig variasjon.')

# 12 R2 svarfordeling
s = content_slide('R2: Svarene spriker mer – flere er fornøyde, men også flere misfornøyde',
                  'Andel som er uenig (1–2), nøytral (3–4) og enig (5–6), sommer 2025 og okt. 2026')
selq = [7, 6, 10, 8, 9]
cats = []; un = []; ne = []; en = []
for i in selq:
    for lab, a in (('2025', a25[i]), ('2026', a26[i])):
        cats.append(f'{Q2[i]} {lab}')
        un.append(a['uenig']); en.append(a['enig']); ne.append(1 - a['uenig'] - a['enig'])
bar_chart(s, 0.3, 1.15, 6.3, 4.25, cats, [('Uenig (1–2)', un), ('Nøytral (3–4)', ne), ('Enig (5–6)', en)],
          [ORANGE, CREAM, DARK], stacked=True, gap=30, size=8, numfmt='0%', label_color=DARK)
tb(s, 6.8, 1.25, 2.75, 4.0, [
    {'text': 'Hva tallene viser', 'size': 11, 'bold': True},
    {'text': 'Enkle å bruke: uenige økte fra 1 % til 17 %, og spredningen (standardavviket) økte fra 1,0 til 1,4.', 'size': 10, 'bullet': True},
    {'text': 'Arbeide effektivt: enige økte fra 24 % til 37 %, men uenige økte også, fra 15 % til 22 %. Snittet er derfor uendret.', 'size': 10, 'bullet': True},
    {'text': 'Frister og prioritering: tydelig flere enige og færre uenige.', 'size': 10, 'bullet': True},
    {'text': 'Mulig forklaring: de som bruker DELE mye og har sakene sine der, får gevinst. Sporadiske brukere og fagområder uten alle sakstyper i DELE opplever omveier.', 'size': 10, 'italic': True},
], margin=0)
notes(s, 'Et uendret snitt kan skjule at gruppen har delt seg. Det gjelder særlig «arbeide effektivt». '
         'Fritekst støtter tolkningen: «For de av oss som bruker systemet i mindre grad er det veldig uoversiktlig» og «mange spørsmål hadde fått høyere score hvis alle relevante sakstyper var inne i DELE».')
# fjern dataetiketter for 2026-rader? beholdes

# 13 R2 fagområder
s = dark_slide('R2: Utvikling per fagområde', 'Måleindikatorer, snitt 2025 → 2026. Grønn = +0,3 eller mer, oransje = −0,3 eller mer. Grå = færre enn 5 svar på indikatoren.')
G = ['Regulatorisk produktinformasjon', 'Regulatorisk før MT', 'Regulatorisk etter MT', 'Sikkerhet – HUM', 'Effekt – HUM',
     'Kvalitet (kjemisk/biologisk)', 'Preklinikk, farmakologi og VET', 'Alle']
MIS2 = [('Medarbeidertilfredshet', 'Medarb.-tilfredshet', list(range(5, 12))), ('Kvalitet', 'Kvalitet', [1, 2]),
        ('Samhandling nasjonal', 'Ekst. nasjonal', [3]), ('Samhandling internasjonal', 'Ekst. internasjonal', [4]),
        ('Informasjonstilgjengelighet', 'Info eksterne', [0])]
rows = [['Fagområde', 'n 25 / 26'] + [m[1] for m in MIS2]]
fills = [[None] * 7]; fgs = [[None] * 7]
for g in G:
    n25 = R2['vals']['2025'].get(g + '_n', R2['n']['2025']); n26 = R2['vals']['2026'].get(g + '_n', R2['n']['2026'])
    row = [g, f'{n25} / {n26}']; fr = [None, None]; fg = [None, None]
    for key, lab, idx in MIS2:
        v25 = mi2['2025'][g].get(key); v26 = mi2['2026'][g].get(key)
        minn = min(min((R2['vals'][y][g][i] or {'n': 0})['n'] for i in idx) for y in ('2025', '2026'))
        row.append(f'{f1(v25)} → {f1(v26)}')
        d = v26 - v25
        if minn < 5:
            fr.append(RGBColor(0x4A, 0x50, 0x3A)); fg.append(RGBColor(0xB0, 0xB0, 0x98))
        else:
            fr.append(GREEN if d >= 0.3 else (ORANGE if d <= -0.3 else CREAM)); fg.append(DARK)
    rows.append(row); fills.append(fr); fgs.append(fg)
cw = [2.35, 0.75] + [1.19] * 5
table(s, 0.47, 1.2, sum(cw), cw, rows, row_h=0.36, size=9, body_fill=None, fills=fills, fgs=fgs, bold_rows=(8,),
      aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 6)
tb(s, 0.47, 4.55, 9.06, 0.8, [
    {'text': 'Sikkerhet – HUM og Preklinikk/farmakologi/VET har mest fremgang. Regulatorisk etter MT har gått ned på alle indikatorer. '
             'Effekt – HUM har lavest nivå. Regulatorisk produktinformasjon er stabil, men «enkle å bruke» falt der fra 4,9 til 3,8.', 'size': 10, 'color': CREAM},
], margin=0)
notes(s, 'Den som har krysset av for flere fagområder, telles i hvert av dem. Kvalitet kjemisk og biologisk er slått sammen, og Preklinikk, farmakologi og VET er slått sammen, for å få minst 5 svar per gruppe. '
         'Regulatorisk etter MT: snittet for å arbeide effektivt falt fra 3,6 til 3,0 og for enkle å bruke fra 3,9 til 3,3 (n 13 → 7). Dette bør følges opp i dialog med fagområdet. '
         'Effekt – HUM: lavest på frister (2,9) og prioritering (2,7). Fritekst herfra peker på at presaker mangler i DELE og at tidslinjer ikke stemmer.')

# 14 Ledere vs saksbehandlere
s = content_slide('Enhetsledere opplever mindre oversikt og styringsstøtte',
                  f'Okt. 2026, R1 og R2 samlet: enhetsledere (n = {nL}) og saksbehandlere/lagledere (n = {nS}). Snitt per spørsmål')
order = list(range(10))
bar_chart(s, 0.3, 1.15, 5.8, 4.25, Q1, [('Saksbehandler/lagleder', mS), ('Enhetsleder', mL)], [LIGHTOLIVE, PURPLE], vmin=0, vmax=6, gap=40, size=9)
tb(s, 6.3, 1.25, 3.25, 4.1, [
    {'text': 'Forskjellen gjelder styring', 'size': 11, 'bold': True},
    {'text': f'Prioritering/ressursstyring: {f1(mL[6])} mot {f1(mS[6])}', 'size': 10, 'bullet': True},
    {'text': f'Status i saksbehandlingen: {f1(mL[3])} mot {f1(mS[3])}', 'size': 10, 'bullet': True},
    {'text': f'Oversikt over oppgaver: {f1(mL[9])} mot {f1(mS[9])}', 'size': 10, 'bullet': True},
    {'text': f'Oversikt over frister: {f1(mL[8])} mot {f1(mS[8])}', 'size': 10, 'bullet': True},
    {'text': 'Lederne trenger rapporter og oversikt på tvers av saker, og det får de ikke i dag. Det sier fritekstsvarene også:', 'size': 10, 'space_after': 6},
    {'text': '«Trenger mer opplæring om alle muligheter som finnes for å få ut rapporter, statistikker og oversikter over ansvarsområder» – enhetsleder, R2', 'size': 9, 'italic': True},
], margin=0)
notes(s, 'Bare 5 enhetsledere svarte (4 i R2, 1 i R1), så tallene er usikre og vises bare samlet av hensyn til anonymitet. '
         'Mønsteret er likevel tydelig og samsvarer med fritekst fra begge releaser (dashbord/PowerBI i R1, rapporter og opplæring i R2). '
         'Lederne er de som skal ta ut gevinsten «tid til koordinering, ressursfordeling og rapportering».')

# =====================================================================
# 15 Section
s = section('Deloverskrift mønster', 'På tvers', 'R1 og R2: tendenser, fritekst, måling og tiltak')
notes(s, 'Tendenser på tvers, forholdet mellom tall og fritekst, og anbefalte tiltak.')

# 16 Tendenser på tvers
s = content_slide('Samme svake punkter i begge releaser',
                  'Snitt per spørsmål: R1 okt. 2026 (alle svar) og R2 sommer 2025 og okt. 2026')
r1all = R1['vals']['2026']['Alle']
bar_chart(s, 0.3, 1.15, 5.8, 4.25, Q1, [('R1 okt. 2026', r1all), ('R2 sommer 2025', [a25[k]['m'] for k in R1MAP]), ('R2 okt. 2026', [a26[k]['m'] for k in R1MAP])],
          [PURPLE, LIGHTOLIVE, DARK], vmin=0, vmax=6, gap=30, overlap=-5, size=9, labels=False)
tb(s, 6.3, 1.25, 3.25, 4.1, [
    {'text': 'Tendenser', 'size': 11, 'bold': True},
    {'text': 'Prioritering/ressursstyring og frister er de svakeste punktene i begge releaser. R2 har bedret seg her, R1 har ikke.', 'size': 10, 'bullet': True},
    {'text': 'Ekstern samhandling og å finne faglig informasjon har ikke bedret seg nevneverdig noe sted. DELE har trolig lite funksjonalitet for dette.', 'size': 10, 'bullet': True},
    {'text': 'Høyest er oversikt over egne oppgaver og status. Dette er det DELE gir mest av i dag.', 'size': 10, 'bullet': True},
    {'text': 'R1-brukerne ligger 0,7–1,7 poeng lavere enn R2-brukerne på alle spørsmål.', 'size': 10, 'bullet': True},
], margin=0)
notes(s, 'R1 og R2 har ulike brukere, så sammenligningen viser mønster, ikke årsak. At R1 (i hovedsak Tilsyn) ligger lavt også etter R2, tyder på at Tilsyns arbeidsprosess (inspeksjoner) ikke er godt nok støttet i DELE.')

# 17 Fritekst vs tall
s = dark_slide('Fritekst og tall forteller samme historie', 'Temaer i fritekst, okt. 2026 (R2: 18 kommentarer, R1: 4 kommentarer)')
rows = [['Tema i fritekst', 'Antall', 'Samsvar med tallene', 'Eksempel'],
        ['Parallelle systemer – ikke alle saker, prosedyrer og presaker i DELE', 'R2: 5', 'Forklarer at effektivitet og faglig støtte står stille, og at svarene spriker.', '«Mye foregår utenfor i en parallell verden med Excel, OneNote, Teams-kanaler og e-post.»'],
        ['Tungvint brukergrensesnitt og mange klikk', 'R2: 2', '«Enkle å bruke» faller (uenige 1 % → 17 %). De som skrev dette, har snitt 1,6 og 3,6.', '«Vanlige arbeidsoppgaver krever unødvendig mange klikk.»'],
        ['Frister og tidslinjer stemmer ikke eller følges utenfor DELE', 'R2: 1, R1: 1', 'Frister er lavest i R1 (Tilsyn 1,9) og i Effekt – HUM (2,9).', '«Frister i DELE er ofte ikke korrekte.»'],
        ['Tilgang til faglig dokumentasjon og litteratur', 'R2: 3', '«Finne faglig info» er uendret (3,6). Gjelder også andre systemer enn DELE.', '«Vanskeligere å finne riktig dokumentasjon nå som Kastor ikke brukes lenger.»'],
        ['Ledere: rapporter, dashbord og opplæring', 'R2: 1, R1: 1', 'Enhetsledere skårer lavt på prioritering og oversikt.', '«PowerBI kunne vært koblet til, bedre dashbord.»'],
        ['Uklart hva undersøkelsen spør om («DELE eller alle verktøy?»)', 'R2: 6', 'Svekker koblingen mellom svarene og DELE.', '«Usikker på om man med IT-verktøy bare mente DELE.»'],
        ]
cw = [2.4, 0.9, 2.85, 2.9]
table(s, 0.47, 1.15, sum(cw), cw, rows, row_h=0.52, size=8, body_fill=None)
tb(s, 0.47, 4.85, 9.06, 0.3, [{'text': 'Positive kommentarer finnes også: «Dele er top», «systemet er et veldig godt utgangspunkt». De som ønsker flere sakstyper i DELE, er ofte fornøyde brukere med høyt snitt (4,5–5,9).',
                              'size': 8, 'color': CREAM, 'italic': True}], margin=0)
notes(s, 'Antallet kommentarer i R2 økte fra 4 (2025) til 18 (2026). Det tyder på økt engasjement rundt DELE. '
         'I R2 har de som kommenterte samme snitt som de andre (3,9), så kommentarene er ikke bare fra misfornøyde. '
         'I R1 er kommentarene mer negative enn snittet (2,2 mot 2,9). '
         'Ønsket om at flere sakstyper skal inn i DELE kommer fra fornøyde brukere. Det er et argument for å utvide bruken, ikke for å endre retning.')

# 18 Måling
s = content_slide('Gevinstmålingen bør forbedres før neste runde')
items = [
    ('Ingen målverdier', 'Uten målverdier kan vi ikke si om en gevinst er realisert. Bør fastsettes for R2 nå, for eksempel snitt ≥ 4,0 eller ≥ 40 % enige.'),
    ('To R1-gevinster er ikke målt', 'Tidsbruk på administrative oppgaver (pri 2.1) og SAM-T-kostnad (pri 1). SAM-T-kostnaden kan trolig dokumenteres raskt fra regnskapet.'),
    ('Uklart hva det spørres om', '6 av 18 kommentarer i R2 handler om undersøkelsen: Gjelder den DELE eller alle verktøy? Spørsmål 12 bør deles i to, og det bør finnes et «vet ikke»-alternativ.'),
    ('Info til eksterne måles på feil sted', '41 % svarer «ikke relevant». Gevinsten gjelder eksterne brukere og bør måles hos dem eller med bruksdata.'),
    ('Rådata og antall svar mangler', 'R1 2022 og 2023 finnes bare som snitt. Rådata og antall svar bør lagres, og skjemaet bør være likt fra gang til gang.'),
]
for i, (h, t) in enumerate(items):
    y = 1.15 + i * 0.83
    rect(s, 0.47, y, 0.55, 0.55, DARK, shape=MSO_SHAPE.OVAL)
    tb(s, 0.47, y, 0.55, 0.55, [{'text': str(i + 1), 'size': 14, 'bold': True, 'color': GREEN, 'align': PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE, margin=0)
    tb(s, 1.2, y - 0.02, 8.3, 0.75, [{'text': h, 'size': 11, 'bold': True, 'space_after': 1}, {'text': t, 'size': 10}], margin=0)
notes(s, 'Dette er forbedringer av selve gevinstoppfølgingen, slik at neste rapport kan si noe om måloppnåelse og ikke bare utvikling.')

# 19 Tiltak
s = dark_slide('Anbefalte tiltak', 'Forslag til eier er veiledende og må avklares i linjen')
rows = [['#', 'Tiltak', 'Adresserer', 'Forslag til eier', 'Når'],
        ['1', 'Eget forbedringsløp for Tilsyn: arbeidsflyt for inspeksjoner, frist- og oppgavestøtte, stabilitet. Avklar om gevinstene for Tilsyn er realistiske.', 'R1 kvalitet, tilfredshet', 'Prosjekteier og leder Tilsyn', 'Q4 2026'],
        ['2', 'Få alle relevante sakstyper, prosedyrer og presaker inn i DELE, og legg ned parallelle lister og verktøy.', 'R2 effektivitet, frister, spredning', 'Produkteier DELE og fagledere', '2027'],
        ['3', 'Gå gjennom brukskvaliteten i de vanligste arbeidsflytene (antall klikk, skjermbilder) og rett feil i frister og tidslinjer.', 'Enkle å bruke, frister', 'Produkteier DELE', 'Q1 2027'],
        ['4', 'Lag rapporter og dashbord for ledere (for eksempel Power BI) og gi opplæring i dem.', 'Ledere: prioritering, oversikt', 'Produkteier DELE og enhetsledere', 'Q1 2027'],
        ['5', 'Avtal felles bruksregler: oppgaver markeres som fullført, og samhandling skjer i DELE. Etabler superbrukere per fagområde.', 'Intern samhandling, status', 'Enhetsledere', 'Løpende'],
        ['6', 'Bedre tilgangen til faglig informasjon (tidsskrifter, databaser, dokumentflyt i Dokubridge og SharePoint).', 'Finne faglig info, kvalitet', 'Linjen og IKT', '2027'],
        ['7', 'Avklar ekstern samhandling: Hvilken funksjonalitet i DELE skal gi gevinsten? Hvis ingen, bør gevinsten tas ut av planen eller flyttes.', 'Ekstern samhandling R1/R2', 'Gevinsteier', 'Q4 2026'],
        ]
cw = [0.3, 4.6, 1.75, 1.6, 0.8]
table(s, 0.47, 1.15, sum(cw), cw, rows, row_h=0.5, size=8, body_fill=None)
notes(s, 'Tiltak 1 og 2 er viktigst for gevinstrealiseringen. Tiltak 1 fordi R1-gevinstene ikke realiseres i Tilsyn. '
         'Tiltak 2 fordi DELE ikke kan gi full gevinst så lenge saksbehandlingen er delt mellom DELE og andre verktøy. '
         'Tiltak 7 handler om realisme i gevinstplanen.')

# 20 Beslutninger
s = content_slide('Forslag til beslutninger for prosjekteier og porteføljestyret')
dec = [
    'Rapporten tas til etterretning. R2 viser fremgang på intern styring og oversikt, mens R1-gevinstene ikke er realisert i Tilsyn.',
    'Det etableres et forbedringsløp for Tilsyn (tiltak 1). Gevinstene for R1 vurderes på nytt når tiltaket er gjennomført.',
    'Målverdier for R2-gevinstene fastsettes og legges inn i Prosjektportalen innen utgangen av 2026.',
    'Gevinsteier dokumenterer SAM-T-kostnaden og avklarer om tidsmåling (R1 pri 2.1) skal gjennomføres eller tas ut av planen.',
    'Gevinsten «ekstern samhandling» knyttes til konkret funksjonalitet i DELE, eller tas ut av planen.',
    'Neste måling gjennomføres høsten 2027 med et forbedret spørreskjema og tidsmåling for utvalgte oppgaver.',
]
for i, d in enumerate(dec):
    y = 1.15 + i * 0.68
    rect(s, 0.47, y, 9.06, 0.58, CREAM if i % 2 == 0 else CYAN, name=f'Beslutning {i+1}')
    tb(s, 0.6, y, 0.4, 0.58, [{'text': str(i + 1), 'size': 16, 'bold': True, 'color': ORANGE}], anchor=MSO_ANCHOR.MIDDLE, margin=0)
    tb(s, 1.0, y, 8.4, 0.58, [{'text': d, 'size': 11}], anchor=MSO_ANCHOR.MIDDLE, margin=0)
notes(s, 'Forslagene kan legges frem som de er, eller justeres før porteføljestyret.')

# =====================================================================
# 21 Vedlegg R2-tabell
s = dark_slide('Vedlegg: R2 per spørsmål', 'n, snitt, andel enig (5–6) og uenig (1–2), «ikke relevant» og p-verdi for endringen')
rows = [['Spørsmål', 'n 25', 'Snitt 25', 'Enig 25', 'Uenig 25', 'n 26', 'Snitt 26', 'Enig 26', 'Uenig 26', 'Ikke rel. 25/26', 'p']]
for i, q in enumerate(Q2):
    a, b = a25[i], a26[i]
    rows.append([q, a['n'], f1(a['m']), pct(a['enig']), pct(a['uenig']), b['n'], f1(b['m']), pct(b['enig']), pct(b['uenig']),
                 f'{a["ir"]} / {b["ir"]}', f'{R2["p"][i]:.2f}'.replace('.', ',')])
fills = [[None] * 11] + [[None] * 10 + [GREEN if R2['p'][i] < 0.05 else None] for i in range(12)]
fgs = [[None] * 11] + [[None] * 10 + [DARK if R2['p'][i] < 0.05 else None] for i in range(12)]
cw = [2.6] + [0.58] * 8 + [1.0, 0.5]
table(s, 0.47, 1.15, sum(cw), cw, rows, row_h=0.3, size=8, body_fill=None, fills=fills, fgs=fgs,
      aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 10)
notes(s, 'p-verdier fra en permutasjonstest (20 000 permutasjoner) av forskjellen i snitt mellom 2025 og 2026. Grønn = p < 0,05.')

# 22 Vedlegg metode / spørsmål
s = content_slide('Vedlegg: Spørsmål og måleindikatorer')
qmap = [
    ('Informasjonstilgjengelighet (kun R2)', 'Vi har gode verktøy for å gjøre informasjon til eksterne brukere enklere tilgjengelig'),
    ('Kvalitet på faglige vurderinger', 'Verktøyene gir støtte til faglige vurderinger · Enkelt å finne detaljert faglig informasjon'),
    ('Samhandling eksterne, nasjonale', 'Gode verktøy for samhandling med eksterne, nasjonale samarbeidspartnere'),
    ('Samhandling eksterne, internasjonale (kun R2)', 'Gode verktøy for samhandling med eksterne, internasjonale samarbeidspartnere'),
    ('Medarbeidertilfredshet', 'Status i saksbehandlingen · Arbeide effektivt · Enkle å bruke · Prioritering og ressursstyring · Intern samhandling · Oversikt over frister · Oversikt over oppgaver'),
]
rows = [['Måleindikator', 'Spørsmål (kortnavn i rapporten)']] + [list(r) for r in qmap]
table(s, 0.47, 1.2, 9.06, [2.8, 6.26], rows, row_h=0.5, size=9, header_fill=DARK, header_fg=CREAM, body_fill=CREAM, body_fg=DARK)
tb(s, 0.47, 4.4, 9.06, 0.9, [
    {'text': 'Grupper i R2: Regulatorisk produktinformasjon, Regulatorisk før MT, Regulatorisk etter MT, Sikkerhet – HUM og Effekt – HUM vises hver for seg. '
             'Kvalitet kjemisk og biologisk er slått sammen. Preklinikk, Farmakologi BE/AM, Farmakologi klinisk og Effekt og sikkerhet VET er slått sammen. '
             'Den som har krysset av for flere fagområder, telles i hvert av dem.', 'size': 9},
], margin=0)
notes(s, 'Metodevedlegg.')

prs.save(OUT)
print('saved', OUT, len(prs.slides))
